"""
formula_generator.py
Modul untuk merender rumus matematika menjadi:
1. Native Office Open XML Math (OMML / <m:oMath>) langsung ke dalam dokumen Microsoft Word (.docx).
   Persamaan menjadi objek native Word Equation yang dapat diedit, tajam (vektor), dan berukuran ringan.
2. Gambar beresolusi tinggi (PNG 300 DPI) di direktori /asset sebagai aset pelengkap atau fallback.
Lengkap dengan penomoran resmi (X.Y) dalam tabel borderless dan keterangan variabel ("di mana: ...").
"""
import os
import re
import xml.sax.saxutils as saxutils
from typing import Dict, Any, List, Optional, Union
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import docx

DEFAULT_ASSET_DIR = "asset"

GREEK_SYMBOLS = {
    r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ",
    r"\epsilon": "ε", r"\varepsilon": "ε", r"\zeta": "ζ", r"\eta": "η",
    r"\theta": "θ", r"\vartheta": "ϑ", r"\iota": "ι", r"\kappa": "κ",
    r"\lambda": "λ", r"\mu": "μ", r"\nu": "ν", r"\xi": "ξ",
    r"\pi": "π", r"\rho": "ρ", r"\sigma": "σ", r"\tau": "τ",
    r"\upsilon": "υ", r"\phi": "φ", r"\chi": "χ", r"\psi": "ψ",
    r"\omega": "ω", r"\Gamma": "Γ", r"\Delta": "Δ", r"\Theta": "Θ",
    r"\Lambda": "Λ", r"\Xi": "Ξ", r"\Pi": "Π", r"\Sigma": "Σ",
    r"\Upsilon": "Υ", r"\Phi": "Φ", r"\Psi": "Ψ", r"\Omega": "Ω"
}

MATH_OPERATORS = {
    r"\times": "×", r"\cdot": "·", r"\pm": "±", r"\mp": "∓",
    r"\div": "÷", r"\leq": "≤", r"\le": "≤", r"\geq": "≥",
    r"\ge": "≥", r"\neq": "≠", r"\approx": "≈", r"\equiv": "≡",
    r"\sim": "∼", r"\in": "∈", r"\notin": "∉", r"\subset": "⊂",
    r"\subseteq": "⊆", r"\rightarrow": "→", r"\leftarrow": "←",
    r"\Rightarrow": "⇒", r"\Leftarrow": "⇐", r"\leftrightarrow": "↔",
    r"\infty": "∞", r"\partial": "∂", r"\nabla": "∇", r"\forall": "∀",
    r"\exists": "∃", r"\ldots": "…", r"\cdots": "⋯", r"\vdots": "⋮"
}

MATH_FUNCTIONS = [
    r"\sin", r"\cos", r"\tan", r"\arcsin", r"\arccos", r"\arctan",
    r"\sinh", r"\cosh", r"\tanh", r"\exp", r"\log", r"\ln",
    r"\min", r"\max", r"\lim", r"\det", r"\dim", r"\gcd"
]

def ensure_asset_dir(workspace_dir: str, asset_folder: str = DEFAULT_ASSET_DIR) -> str:
    """Memastikan folder asset dibuat di dalam workspace."""
    target_path = os.path.join(workspace_dir, asset_folder)
    os.makedirs(target_path, exist_ok=True)
    return target_path

def clean_latex_string(latex_code: str) -> str:
    """Membersihkan dan menormalisasi string LaTeX."""
    s = latex_code.strip()
    if s.startswith("$$") and s.endswith("$$"):
        s = s[2:-2].strip()
    elif s.startswith("$") and s.endswith("$"):
        s = s[1:-1].strip()
    return f"${s}$"

def extract_curly_group(s: str, start_idx: int):
    """Mengambil isi dalam kurung kurawal { ... } yang seimbang."""
    if start_idx >= len(s) or s[start_idx] != '{':
        return "", start_idx
    depth = 0
    content = []
    i = start_idx
    while i < len(s):
        c = s[i]
        if c == '{':
            depth += 1
            if depth > 1:
                content.append(c)
        elif c == '}':
            depth -= 1
            if depth == 0:
                return "".join(content), i + 1
            else:
                content.append(c)
        else:
            content.append(c)
        i += 1
    return "".join(content), i

def parse_latex_to_omml_xml(latex_str: str) -> str:
    """
    Mengonversi notasi matematika LaTeX ke struktur XML Office Open XML Math (OMML).
    Dihasilkan sebagai elemen <m:oMath> yang siap disematkan langsung ke python-docx.
    """
    s = latex_str.strip()
    if s.startswith("$$") and s.endswith("$$"):
        s = s[2:-2].strip()
    elif s.startswith("$") and s.endswith("$"):
        s = s[1:-1].strip()

    s = s.replace(r"\displaystyle", "").replace(r"\textstyle", "").strip()
    inner_xml = _convert_tokens_to_omml(s)
    return f'<m:oMath {nsdecls("m")}>{inner_xml}</m:oMath>'

def _convert_tokens_to_omml(latex: str) -> str:
    xml_parts = []
    i = 0
    n = len(latex)

    while i < n:
        # 1. Whitespace
        if latex[i].isspace():
            i += 1
            continue

        # 2. Fractions: \frac{num}{den}
        if latex[i:].startswith(r"\frac"):
            i += 5
            while i < n and latex[i].isspace():
                i += 1
            num, i = extract_curly_group(latex, i)
            while i < n and latex[i].isspace():
                i += 1
            den, i = extract_curly_group(latex, i)

            num_xml = _convert_tokens_to_omml(num)
            den_xml = _convert_tokens_to_omml(den)
            xml_parts.append(
                f'<m:f><m:fPr><m:type m:val="bar"/></m:fPr>'
                f'<m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>'
            )
            continue

        # 3. Square root: \sqrt{arg} or \sqrt[deg]{arg}
        if latex[i:].startswith(r"\sqrt"):
            i += 5
            deg_str = ""
            while i < n and latex[i].isspace():
                i += 1
            if i < n and latex[i] == '[':
                deg_end = latex.find(']', i)
                if deg_end != -1:
                    deg_str = latex[i+1:deg_end]
                    i = deg_end + 1
            while i < n and latex[i].isspace():
                i += 1
            arg, i = extract_curly_group(latex, i)
            arg_xml = _convert_tokens_to_omml(arg)
            if deg_str:
                deg_xml = _convert_tokens_to_omml(deg_str)
                xml_parts.append(
                    f'<m:rad><m:radPr><m:degHide m:val="0"/></m:radPr>'
                    f'<m:deg>{deg_xml}</m:deg><m:e>{arg_xml}</m:e></m:rad>'
                )
            else:
                xml_parts.append(
                    f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
                    f'<m:deg/><m:e>{arg_xml}</m:e></m:rad>'
                )
            continue

        # 4. Text labels: \text{...}, \mathrm{...}, \mathbf{...}
        m_text = re.match(r'\\(?:text|mathrm|mathbf|mathit)\s*\{', latex[i:])
        if m_text:
            start_bracket = i + m_text.end() - 1
            txt_content, i = extract_curly_group(latex, start_bracket)
            esc_txt = saxutils.escape(txt_content)
            xml_parts.append(f'<m:r><m:rPr><m:nor/></m:rPr><m:t>{esc_txt}</m:t></m:r>')
            continue

        # 5. Standard Math Functions (\sin, \cos, \log, \exp, \lim, dll)
        matched_fn = False
        for fn in MATH_FUNCTIONS:
            if latex[i:].startswith(fn):
                tail = latex[i+len(fn):]
                if not tail or not tail[0].isalpha():
                    fn_name = fn[1:]
                    xml_parts.append(f'<m:r><m:rPr><m:nor/></m:rPr><m:t>{fn_name}</m:t></m:r>')
                    i += len(fn)
                    matched_fn = True
                    break
        if matched_fn:
            continue

        # 6. N-ary Summation, Product, Integral (\sum, \prod, \int)
        m_nary = re.match(r'\\(sum|prod|int|oint|bigcap|bigcup)', latex[i:])
        if m_nary:
            cmd = m_nary.group(1)
            i += m_nary.end()
            chr_sym = {"sum": "∑", "prod": "∏", "int": "∫", "oint": "∮", "bigcap": "⋂", "bigcup": "⋃"}.get(cmd, "∑")

            sub_xml = ""
            sup_xml = ""

            while i < n and latex[i].isspace():
                i += 1
            if i < n and latex[i] == '_':
                i += 1
                while i < n and latex[i].isspace():
                    i += 1
                if i < n and latex[i] == '{':
                    sub_val, i = extract_curly_group(latex, i)
                else:
                    sub_val = latex[i] if i < n else ""
                    i += 1
                sub_xml = _convert_tokens_to_omml(sub_val)

            while i < n and latex[i].isspace():
                i += 1
            if i < n and latex[i] == '^':
                i += 1
                while i < n and latex[i].isspace():
                    i += 1
                if i < n and latex[i] == '{':
                    sup_val, i = extract_curly_group(latex, i)
                else:
                    sup_val = latex[i] if i < n else ""
                    i += 1
                sup_xml = _convert_tokens_to_omml(sup_val)

            xml_parts.append(
                f'<m:nary><m:naryPr><m:chr m:val="{chr_sym}"/><m:limLoc m:val="undOvr"/></m:naryPr>'
                f'<m:sub>{sub_xml}</m:sub><m:sup>{sup_xml}</m:sup><m:e/></m:nary>'
            )
            continue

        # 7. Delimiters: \left( ... \right), | ... |
        if latex[i:].startswith(r"\left"):
            i += 5
            beg_chr = "("
            if i < n:
                if latex[i] == '[': beg_chr = "["
                elif latex[i] == '{' or latex[i:].startswith(r"\{"): beg_chr = "{"
                elif latex[i] == '|': beg_chr = "|"
                elif latex[i] == '(': beg_chr = "("
                i += 1

            right_idx = latex.find(r"\right", i)
            if right_idx != -1:
                inner_content = latex[i:right_idx]
                i = right_idx + 6
                end_chr = ")"
                if i < n:
                    if latex[i] == ']': end_chr = "]"
                    elif latex[i] == '}' or latex[i:].startswith(r"\}"): end_chr = "}"
                    elif latex[i] == '|': end_chr = "|"
                    elif latex[i] == ')': end_chr = ")"
                    i += 1
                inner_d_xml = _convert_tokens_to_omml(inner_content)
                xml_parts.append(
                    f'<m:d><m:dPr><m:begChr m:val="{beg_chr}"/><m:endChr m:val="{end_chr}"/></m:dPr>'
                    f'<m:e>{inner_d_xml}</m:e></m:d>'
                )
                continue

        # 8. Accents: \hat{x}, \bar{x}, \tilde{x}, \vec{x}
        m_acc = re.match(r'\\(hat|bar|tilde|vec|dot|ddot)\s*\{', latex[i:])
        if m_acc:
            acc_name = m_acc.group(1)
            acc_chr = {"hat": "^", "bar": "¯", "tilde": "~", "vec": "→", "dot": "˙", "ddot": "¨"}.get(acc_name, "^")
            start_bracket = i + m_acc.end() - 1
            inner_acc, i = extract_curly_group(latex, start_bracket)
            inner_acc_xml = _convert_tokens_to_omml(inner_acc)
            xml_parts.append(
                f'<m:acc><m:accPr><m:chr m:val="{acc_chr}"/></m:accPr>'
                f'<m:e>{inner_acc_xml}</m:e></m:acc>'
            )
            continue

        # 9. Paired Pipe Delimiters | ... |
        if latex[i] == '|':
            next_pipe = latex.find('|', i + 1)
            if next_pipe != -1 and '\n' not in latex[i+1:next_pipe]:
                inner_pipe = latex[i+1:next_pipe]
                i = next_pipe + 1
                inner_d_xml = _convert_tokens_to_omml(inner_pipe)
                xml_parts.append(
                    f'<m:d><m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr>'
                    f'<m:e>{inner_d_xml}</m:e></m:d>'
                )
                continue

        # 8. Greek symbols & math operators
        matched_sym = False
        for tex_k, sym_v in {**GREEK_SYMBOLS, **MATH_OPERATORS}.items():
            if latex[i:].startswith(tex_k):
                tail = latex[i+len(tex_k):]
                if not tail or not tail[0].isalpha():
                    xml_parts.append(f'<m:r><m:t>{sym_v}</m:t></m:r>')
                    i += len(tex_k)
                    matched_sym = True
                    break
        if matched_sym:
            continue

        # 9. Subscripts and Superscripts on regular tokens
        m_base = re.match(r'([a-zA-Z0-9]+)', latex[i:])
        if m_base:
            base_txt = m_base.group(1)
            i += m_base.end()

            has_sub = False
            has_sup = False
            sub_val = ""
            sup_val = ""

            while i < n and (latex[i] in ['_', '^'] or latex[i].isspace()):
                if latex[i].isspace():
                    i += 1
                    continue
                if latex[i] == '_':
                    has_sub = True
                    i += 1
                    while i < n and latex[i].isspace(): i += 1
                    if i < n and latex[i] == '{':
                        sub_val, i = extract_curly_group(latex, i)
                    else:
                        sub_val = latex[i] if i < n else ""
                        i += 1
                elif latex[i] == '^':
                    has_sup = True
                    i += 1
                    while i < n and latex[i].isspace(): i += 1
                    if i < n and latex[i] == '{':
                        sup_val, i = extract_curly_group(latex, i)
                    else:
                        sup_val = latex[i] if i < n else ""
                        i += 1

            base_xml = f'<m:r><m:t>{saxutils.escape(base_txt)}</m:t></m:r>'
            if has_sub and has_sup:
                sub_xml = _convert_tokens_to_omml(sub_val)
                sup_xml = _convert_tokens_to_omml(sup_val)
                xml_parts.append(f'<m:sSubSup><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub><m:sup>{sup_xml}</m:sup></m:sSubSup>')
            elif has_sub:
                sub_xml = _convert_tokens_to_omml(sub_val)
                xml_parts.append(f'<m:sSub><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub></m:sSub>')
            elif has_sup:
                sup_xml = _convert_tokens_to_omml(sup_val)
                xml_parts.append(f'<m:sSup><m:e>{base_xml}</m:e><m:sup>{sup_xml}</m:sup></m:sSup>')
            else:
                xml_parts.append(base_xml)
            continue

        # 10. Default single char / operator
        c = latex[i]
        esc_c = saxutils.escape(c)
        if c in ['=', '+', '-', '*', '/', '<', '>']:
            xml_parts.append(f'<m:r><m:t> {esc_c} </m:t></m:r>')
        else:
            xml_parts.append(f'<m:r><m:t>{esc_c}</m:t></m:r>')
        i += 1

    return "".join(xml_parts)

def render_latex_formula_image(
    latex_code: str,
    output_path: str,
    fontsize: int = 14,
    dpi: int = 300,
    text_color: str = "#0F172A"
) -> str:
    """Merender rumus matematika LaTeX menjadi gambar PNG transparan beresolusi tinggi (300 DPI)."""
    formatted_latex = clean_latex_string(latex_code)
    fig = plt.figure(figsize=(7.0, 1.2), dpi=dpi)
    fig.patch.set_alpha(0.0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis('off')
    ax.patch.set_alpha(0.0)

    try:
        ax.text(0.5, 0.5, formatted_latex, fontsize=fontsize, ha='center', va='center', color=text_color)
    except Exception:
        plain_text = latex_code.replace("$", "").strip()
        ax.text(0.5, 0.5, plain_text, fontsize=fontsize - 2, ha='center', va='center', fontfamily='serif', fontstyle='italic', color=text_color)

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    plt.savefig(output_path, dpi=dpi, bbox_inches='tight', transparent=True, pad_inches=0.08)
    plt.close(fig)
    return output_path

def add_formula_to_docx_doc(
    doc: docx.Document,
    latex_code: Optional[str] = None,
    image_path: Optional[str] = None,
    chapter_num: int = 3,
    formula_num: int = 1,
    intro_text: Optional[str] = None,
    variable_definitions: Optional[Dict[str, str]] = None,
    use_native_equation: bool = True,
    image_width_inches: float = 4.0
):
    """
    Menyisipkan rumus matematika ke dokumen DOCX dengan format resmi:
    1. Paragraf pengantar (opsional).
    2. Tabel 1 baris x 2 kolom borderless:
       - Kolom kiri (5.2 inci): Objek Persamaan Native Word (<m:oMath>) terpusat (Center).
       - Kolom kanan (0.8 inci): Nomor persamaan (X.Y) rata kanan (Right).
    3. Keterangan simbol variabel 'di mana:' (opsional).
    """
    if intro_text:
        p_intro = doc.add_paragraph(intro_text)
        p_intro.paragraph_format.space_after = Pt(4)

    # 1. Buat tabel borderless 1x2 untuk tata letak rumus & nomor persamaan
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Hilangkan border tabel
    tblPr = tbl._tbl.tblPr
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>'
        f'<w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

    # Kolom 0: Rumus Matematika (Center)
    cell_eq = tbl.rows[0].cells[0]
    cell_eq.width = Inches(5.2)
    p_eq = cell_eq.paragraphs[0]
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(4)
    p_eq.paragraph_format.space_after = Pt(4)

    inserted_native = False
    if use_native_equation and latex_code:
        try:
            omml_xml = parse_latex_to_omml_xml(latex_code)
            omml_elem = parse_xml(omml_xml)
            p_eq._p.append(omml_elem)
            inserted_native = True
        except Exception:
            inserted_native = False

    if not inserted_native:
        # Fallback ke render citra jika native gagal atau diminta citra
        if image_path and os.path.exists(image_path):
            run_img = p_eq.add_run()
            run_img.add_picture(image_path, width=Inches(image_width_inches))
        elif latex_code:
            run_fallback = p_eq.add_run(latex_code.replace("$", ""))
            run_fallback.italic = True

    # Kolom 1: Nomor Persamaan (Right)
    cell_num = tbl.rows[0].cells[1]
    cell_num.width = Inches(0.8)
    p_num = cell_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(8)
    p_num.paragraph_format.space_after = Pt(4)
    run_num = p_num.add_run(f"({chapter_num}.{formula_num})")
    run_num.bold = False

    # 2. Tambahkan keterangan variabel jika ada
    if variable_definitions:
        p_dimana = doc.add_paragraph("di mana:")
        p_dimana.paragraph_format.space_before = Pt(4)
        p_dimana.paragraph_format.space_after = Pt(2)

        for sym, desc in variable_definitions.items():
            p_var = doc.add_paragraph()
            p_var.paragraph_format.left_indent = Inches(0.25)
            p_var.paragraph_format.space_before = Pt(1)
            p_var.paragraph_format.space_after = Pt(2)

            r_sym = p_var.add_run(f"{sym}")
            r_sym.italic = True
            p_var.add_run(f" = {desc}")

def generate_math_formula(
    latex_code: str,
    formula_title: str = "Persamaan Matematika",
    chapter_num: int = 3,
    formula_num: int = 1,
    variable_definitions: Optional[Dict[str, str]] = None,
    workspace_dir: str = ".",
    asset_folder: str = DEFAULT_ASSET_DIR,
    output_filename: Optional[str] = None,
    target_document_docx: Optional[str] = None,
    intro_text: Optional[str] = None,
    use_native_equation: bool = True,
    save_png_asset: bool = True
) -> Dict[str, Any]:
    """
    Fungsi orkestrasi untuk membuat rumus matematika native Word (OMML)
    dan/atau gambar PNG di /asset, serta menyisipkannya ke dokumen DOCX.
    """
    asset_dir = ensure_asset_dir(workspace_dir=workspace_dir, asset_folder=asset_folder)

    if not output_filename:
        safe_title = re.sub(r'[^a-zA-Z0-9_-]', '_', formula_title.lower())[:25]
        output_filename = f"formula_{chapter_num}_{formula_num}_{safe_title}.png"

    img_path = os.path.join(asset_dir, output_filename)
    rendered_file = None
    if save_png_asset:
        rendered_file = render_latex_formula_image(
            latex_code=latex_code,
            output_path=img_path
        )

    relative_path = os.path.join(asset_folder, output_filename).replace("\\", "/") if save_png_asset else None
    inserted_info = None

    if target_document_docx:
        doc_path = os.path.join(workspace_dir, target_document_docx)
        if os.path.exists(doc_path):
            doc = docx.Document(doc_path)
            add_formula_to_docx_doc(
                doc=doc,
                latex_code=latex_code,
                image_path=img_path if save_png_asset else None,
                chapter_num=chapter_num,
                formula_num=formula_num,
                intro_text=intro_text or f"Persamaan {formula_title} dirumuskan sebagai berikut:",
                variable_definitions=variable_definitions,
                use_native_equation=use_native_equation
            )
            doc.save(doc_path)
            inserted_info = {
                "document": target_document_docx,
                "equation_number": f"({chapter_num}.{formula_num})",
                "rendering_mode": "NATIVE_OMML_EQUATION" if use_native_equation else "IMAGE_PNG",
                "status": "INSERTED"
            }

    return {
        "status": "SUCCESS",
        "formula_title": formula_title,
        "latex_code": latex_code,
        "rendering_mode": "NATIVE_OMML_EQUATION" if use_native_equation else "IMAGE_PNG",
        "equation_number": f"({chapter_num}.{formula_num})",
        "image_filename": output_filename if save_png_asset else None,
        "asset_folder": asset_folder,
        "relative_path": relative_path,
        "full_path": os.path.abspath(rendered_file) if rendered_file else None,
        "file_size_bytes": os.path.getsize(rendered_file) if (rendered_file and os.path.exists(rendered_file)) else 0,
        "variable_definitions": variable_definitions or {},
        "inserted_into_document": inserted_info
    }
