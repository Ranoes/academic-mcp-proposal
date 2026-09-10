"""
builder_engine.py
Independent Academic Proposal Builder Engine
Kompatibel dengan Template Resmi FILKOM UB v3.0 dan berbagai variasi proposal akademik.
"""
import os
import re
import json
import docx
from datetime import datetime
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from typing import Dict, Any, List, Optional

DEFAULT_TEMPLATE_PATH = "/app/assets/templates/template_filkom_ub_v3.0.docx"
# Fallback lokal jika dijalankan di luar docker
if not os.path.exists(DEFAULT_TEMPLATE_PATH):
    DEFAULT_TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "assets", "templates", "template_filkom_ub_v3.0.docx")

def clean_heading_title(title: str, level: int = 1) -> str:
    """
    Membersihkan penomoran manual dari judul bab/subbab karena style Heading pada
    template resmi FILKOM UB (Heading 1, Heading 2, Heading 3) sudah memiliki
    multilevel numbering otomatis (BAB 1, 1.1, 1.1.1).
    """
    if not title:
        return ""
    t = title.strip()
    if level == 1:
        # Hapus prefix manual seperti "BAB 1: ", "BAB I - ", "BAB 1. "
        t = re.sub(r'^(?:BAB\s+[0-9IVXLCDM]+[\s:\-\.]*)\s*', '', t, flags=re.IGNORECASE)
        return t.upper()
    else:
        # Hapus prefix manual nomor seperti "1.1 ", "2.1.1 ", "1. ", "A.1 "
        t = re.sub(r'^(?:[0-9]+(?:\.[0-9]+)*|[A-Z]\.[0-9]+)[\s:\-\.]+\s*', '', t)
        return t.strip()

def strip_list_prefix(text: str) -> str:
    """
    Menghapus penomoran manual seperti "1. ", "a. ", "- ", "* "
    agar dapat diformat secara bersih oleh style list Word.
    """
    if not text:
        return ""
    t = text.strip()
    return re.sub(r'^(?:\d+[\.\)]|[a-zA-Z][\.\)]|[\-\*\•])\s*', '', t)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    """
    Menerapkan border standar template FILKOM UB: garis tunggal hitam (0.5 pt / 4 eighths).
    """
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_cell_header(cell, text: str):
    """
    Format sel header tabel sesuai style resmi FILKOM UB (Table Column Title, 10pt bold, centered, tanpa shading abu-abu).
    """
    p = cell.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        p.style = 'Table Column Title'
    except Exception:
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

def format_cell_body(cell, text: str, align=WD_ALIGN_PARAGRAPH.LEFT, bold: bool = False):
    """
    Format sel badan tabel sesuai style resmi FILKOM UB (Table Content, 10pt normal).
    """
    p = cell.paragraphs[0]
    p.text = text
    p.alignment = align
    try:
        p.style = 'Table Content'
        if bold:
            for r in p.runs:
                r.font.bold = True
    except Exception:
        for r in p.runs:
            r.font.bold = bold
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)

def apply_document_default_font(doc, font_name: str = "Times New Roman"):
    """
    Mengubah default font seluruh dokumen dan style Word ke font_name (default: Times New Roman).
    """
    styles_elm = doc.styles.element
    # 1. Update docDefaults
    doc_defaults = styles_elm.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docDefaults')
    if doc_defaults is not None:
        rPrDefault = doc_defaults.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPrDefault')
        if rPrDefault is not None:
            rPr = rPrDefault.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
            if rPr is not None:
                rFonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                if rFonts is not None:
                    rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii', font_name)
                    rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi', font_name)
                    rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs', font_name)
                    rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', font_name)

    # 2. Update style fonts
    for s in doc.styles:
        if s.type in (docx.enum.style.WD_STYLE_TYPE.PARAGRAPH, docx.enum.style.WD_STYLE_TYPE.CHARACTER):
            if 'code' in s.name.lower() or 'source' in s.name.lower():
                continue
            if s.font:
                s.font.name = font_name

    # 3. Update rFonts pada level styles.xml
    for rFonts in styles_elm.xpath('//w:rFonts'):
        ascii_font = rFonts.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii')
        if ascii_font and 'courier' in ascii_font.lower():
            continue
        for attr in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            key = f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}'
            if key in rFonts.attrib:
                rFonts.attrib[key] = font_name

def add_table_caption(doc, chapter_num: int, table_num: int, title: str):
    """
    Menambahkan caption tabel sesuai Lampiran A & Subbab 2.2.2 Template FILKOM UB:
    1. Diletakkan di ATAS tabel.
    2. Format: Tabel <STYLEREF 1 \\s>.<SEQ Tabel \\* ARABIC \\s 1> <Judul>
    3. Judul tabel ditulis tebal (bold), diawali huruf kapital, tanpa diakhiri tanda titik.
    4. Menyertakan field Word SEQ Tabel sehingga otomatis terindeks pada DAFTAR TABEL.
    """
    clean_title = title.strip().rstrip('.')
    p = doc.add_paragraph()
    try:
        p.style = 'Caption'
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption_xml = (
        f'<w:r {nsdecls("w")}><w:t xml:space="preserve">Tabel </w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> STYLEREF 1 \\s </w:instrText></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>{chapter_num}</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>.</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> SEQ Tabel \\* ARABIC \\s 1 </w:instrText></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>{table_num}</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t xml:space="preserve"> </w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:rPr><w:b/></w:rPr><w:t>{clean_title}</w:t></w:r>'
    )
    for elem in parse_xml(f'<w:p {nsdecls("w")}>{caption_xml}</w:p>'):
        p._p.append(elem)
    return p

def add_figure_caption(doc, chapter_num: int, figure_num: int, title: str):
    """
    Menambahkan caption gambar sesuai Lampiran A & Subbab 2.2.3 Template FILKOM UB:
    1. Diletakkan di BAWAH gambar.
    2. Format: Gambar <STYLEREF 1 \\s>.<SEQ Gambar \\* ARABIC \\s 1> <Judul>
    3. Judul gambar ditulis tebal (bold), diawali huruf kapital, tanpa diakhiri tanda titik.
    4. Menyertakan field Word SEQ Gambar sehingga otomatis terindeks pada DAFTAR GAMBAR.
    """
    clean_title = title.strip().rstrip('.')
    p = doc.add_paragraph()
    try:
        p.style = 'Caption'
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption_xml = (
        f'<w:r {nsdecls("w")}><w:t xml:space="preserve">Gambar </w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> STYLEREF 1 \\s </w:instrText></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>{chapter_num}</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>.</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> SEQ Gambar \\* ARABIC \\s 1 </w:instrText></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t>{figure_num}</w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'
        f'<w:r {nsdecls("w")}><w:t xml:space="preserve"> </w:t></w:r>'
        f'<w:r {nsdecls("w")}><w:rPr><w:b/></w:rPr><w:t>{clean_title}</w:t></w:r>'
    )
    for elem in parse_xml(f'<w:p {nsdecls("w")}>{caption_xml}</w:p>'):
        p._p.append(elem)
    return p

def add_styled_table(
    doc,
    data: List[List[str]],
    center_cols: Optional[List[int]] = None
):
    """
    Membuat tabel dengan layout resmi FILKOM UB:
    - Alignment: Center
    - Borders: Garis tunggal hitam (0.5 pt)
    - Row 0: Table Column Title, centered, w:tblHeader agar berulang di halaman lanjutan
    - Row 1+: Table Content, left-aligned / centered, w:cantSplit agar baris tidak terpotong halaman
    """
    if not data or not data[0]:
        return None

    num_rows = len(data)
    num_cols = len(data[0])
    center_cols = center_cols or []

    t = doc.add_table(rows=num_rows, cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)

    # Header Row
    hdr_row = t.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for col_idx, text in enumerate(data[0]):
        cell = hdr_row.cells[col_idx]
        format_cell_header(cell, text)

    # Data Rows
    for row_idx, row_data in enumerate(data[1:], start=1):
        row = t.rows[row_idx]
        row_trPr = row._tr.get_or_add_trPr()
        row_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_body(cell, text, align=align)

    return t

def add_figure(doc, image_path: str, chapter_num: int, figure_num: int, title: str, width_inches: float = 5.5):
    """
    Menambahkan gambar dan caption gambar di bawahnya sesuai ketentuan FILKOM UB.
    """
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        return add_figure_caption(doc, chapter_num, figure_num, title)
    return None

def create_generic_proposal(
    metadata: Dict[str, str],
    bab1_data: Dict[str, Any],
    bab2_subbab: List[Dict[str, Any]],
    bab3_subbab: List[Dict[str, Any]],
    daftar_referensi: List[str],
    tabel_tinjauan_pustaka: Optional[List[List[str]]] = None,
    tabel_operasionalisasi_variabel: Optional[List[List[str]]] = None,
    tabel_tahapan_metode: Optional[List[List[str]]] = None,
    tabel_jadwal: Optional[List[List[str]]] = None,
    lampiran_data: Optional[List[Dict[str, Any]]] = None,
    default_font: str = "Times New Roman",
    output_path: str = "Proposal Skripsi.docx",
    template_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Membangun dokumen DOCX proposal lengkap dan rapi dengan mengadopsi format resmi Template FILKOM UB v3.0:
    - Default font: Times New Roman (atau font yang dikustomisasi)
    - Cover page: Title (16pt bold uppercase), Author, Institution
    - Section 0: Halaman romawi (Daftar Isi, Daftar Tabel, Daftar Gambar, Daftar Lampiran)
    - Section 1: Bab 1, Bab 2, Bab 3, Daftar Referensi, Lampiran (angka arab mulai 1, margin 4-3-3-3)
    - Multilevel heading tanpa nomor manual ganda
    - Format tabel resmi: Table Column Title, Table Content, cantSplit, tblHeader, border tunggal
    - Caption tabel di atas tabel (Tabel X.Y Judul) dan caption gambar di bawah gambar
    - Daftar referensi style Harvard dengan hanging indent (style References)
    """
    tpl = template_path or DEFAULT_TEMPLATE_PATH
    if not os.path.exists(tpl):
        raise FileNotFoundError(f"Template berkas tidak ditemukan di path: {tpl}")

    doc = docx.Document(tpl)
    apply_document_default_font(doc, default_font)

    # =========================================================================
    # 1. UPDATE HALAMAN JUDUL (COVER)
    # =========================================================================
    judul = metadata.get("judul", "JUDUL PROPOSAL SKRIPSI").upper()
    nama_mhs = metadata.get("nama_mahasiswa", "Nama Mahasiswa")
    nim_mhs = metadata.get("nim", "NIM Mahasiswa")
    prodi = metadata.get("program_studi", "PROGRAM STUDI SISTEM INFORMASI").upper()
    departemen = (metadata.get("departemen") or metadata.get("jurusan") or "DEPARTEMEN SISTEM INFORMASI").upper()
    fakultas = metadata.get("fakultas", "FAKULTAS ILMU KOMPUTER").upper()
    universitas = metadata.get("universitas", "UNIVERSITAS BRAWIJAYA").upper()
    kota = metadata.get("kota", "MALANG").upper()
    tahun = str(metadata.get("tahun", datetime.now().year))

    doc.paragraphs[0].text = judul
    doc.paragraphs[6].text = nama_mhs
    doc.paragraphs[7].text = f"NIM: {nim_mhs}"
    doc.paragraphs[18].text = prodi
    doc.paragraphs[19].text = departemen
    doc.paragraphs[20].text = fakultas
    doc.paragraphs[21].text = universitas
    doc.paragraphs[22].text = kota
    doc.paragraphs[23].text = tahun

    # =========================================================================
    # 2. HAPUS PANDUAN PLACEHOLDER TEMPLATE GUIDE (INDEX 75 SAMPAI SEBELUM FINAL sectPr)
    # =========================================================================
    # Index 74 adalah paragraf pembatas Section 0 (dengan sectPr romawi)
    # Index 75 s.d. -1 adalah materi panduan buku yang harus digantikan
    body = doc._body._body
    to_remove = body[75:-1]
    for elem in to_remove:
        body.remove(elem)

    # Perbaiki cache nomor halaman awal pada footer Seksi 1 (angka '8' bawaan template diubah ke '1')
    if len(doc.sections) > 1:
        s1 = doc.sections[1]
        for f in [s1.footer, s1.first_page_footer, s1.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for r in p.runs:
                        if r.text.strip() == '8':
                            r.text = '1'

    # Helper penambahan elemen dengan style resmi FILKOM UB
    def add_p(text: str, style='Body Text First Indent'):
        return doc.add_paragraph(text, style=style)

    def add_h1(text: str):
        clean_txt = clean_heading_title(text, level=1)
        return doc.add_paragraph(clean_txt, style='Heading 1')

    def add_h2(text: str):
        clean_txt = clean_heading_title(text, level=2)
        return doc.add_paragraph(clean_txt, style='Heading 2')

    def add_h3(text: str):
        clean_txt = clean_heading_title(text, level=3)
        return doc.add_paragraph(clean_txt, style='Heading 3')

    def add_list_number(text: str):
        clean_txt = strip_list_prefix(text)
        style = 'List Number' if 'List Number' in doc.styles else 'List Paragraph'
        return doc.add_paragraph(clean_txt, style=style)

    def add_list_bullet(text: str):
        clean_txt = strip_list_prefix(text)
        style = 'List Bullet' if 'List Bullet' in doc.styles else 'List Paragraph'
        return doc.add_paragraph(clean_txt, style=style)

    def add_ref(text: str):
        style = 'References' if 'References' in doc.styles else 'Normal'
        return doc.add_paragraph(text, style=style)

    # =========================================================================
    # BAB 1 PENDAHULUAN
    # =========================================================================
    add_h1("PENDAHULUAN")

    add_h2("Latar Belakang")
    for para in bab1_data.get("latar_belakang", []):
        add_p(para)

    add_h2("Rumusan Masalah")
    add_p(bab1_data.get("rumusan_masalah_pengantar", "Berdasarkan konteks latar belakang yang telah diuraikan, rumusan masalah utama dalam penelitian ini adalah:"))
    add_p(bab1_data.get("rumusan_masalah", ""))
    if bab1_data.get("rumusan_masalah_penjelasan"):
        add_p(bab1_data.get("rumusan_masalah_penjelasan"))

    add_h2("Tujuan")
    if bab1_data.get("tujuan_umum"):
        add_p(f"Tujuan umum dari penelitian ini adalah {bab1_data.get('tujuan_umum')}")
    if bab1_data.get("tujuan_khusus"):
        add_p("Tujuan khusus dari penelitian ini dijabarkan sebagai berikut:")
        for tk in bab1_data.get("tujuan_khusus", []):
            add_list_number(tk)

    add_h2("Manfaat")
    add_p("Pelaksanaan penelitian ini diharapkan memberikan manfaat teoritis dan praktis:")
    for mf in bab1_data.get("manfaat", []):
        add_list_number(mf)

    add_h2("Batasan Masalah")
    for bm in bab1_data.get("batasan_masalah", []):
        add_list_number(bm)

    add_h2("Sistematika Pembahasan")
    for sp in bab1_data.get("sistematika_pembahasan", []):
        add_p(sp)

    # =========================================================================
    # BAB 2 LANDASAN KEPUSTAKAAN
    # =========================================================================
    add_h1("LANDASAN KEPUSTAKAAN")
    for section in bab2_subbab:
        lvl = section.get("level", 2)
        title = section.get("title", "")
        if lvl == 2:
            add_h2(title)
        else:
            add_h3(title)
        for para in section.get("paragraphs", []):
            add_p(para)

    if tabel_tinjauan_pustaka:
        add_table_caption(doc, chapter_num=2, table_num=1, title="Matriks Kajian Penelitian Terdahulu")
        add_styled_table(doc, tabel_tinjauan_pustaka, center_cols=[0])

    # =========================================================================
    # BAB 3 METODOLOGI
    # =========================================================================
    add_h1("METODOLOGI")
    for section in bab3_subbab:
        lvl = section.get("level", 2)
        title = section.get("title", "")
        if lvl == 2:
            add_h2(title)
        else:
            add_h3(title)
        for para in section.get("paragraphs", []):
            add_p(para)

    ch3_tbl_idx = 1
    if tabel_operasionalisasi_variabel:
        add_table_caption(doc, chapter_num=3, table_num=ch3_tbl_idx, title="Operasionalisasi Konstruk dan Variabel Penelitian")
        add_styled_table(doc, tabel_operasionalisasi_variabel, center_cols=[0, 4])
        ch3_tbl_idx += 1

    if tabel_tahapan_metode:
        add_table_caption(doc, chapter_num=3, table_num=ch3_tbl_idx, title="Matriks Tahapan Metodologi dan Luaran Artefak")
        add_styled_table(doc, tabel_tahapan_metode, center_cols=[0])
        ch3_tbl_idx += 1

    if tabel_jadwal:
        add_table_caption(doc, chapter_num=3, table_num=ch3_tbl_idx, title="Rencana Jadwal Pelaksanaan Penelitian")
        center_cols_jadwal = list(range(1, len(tabel_jadwal[0]))) if len(tabel_jadwal) > 0 else []
        add_styled_table(doc, tabel_jadwal, center_cols=center_cols_jadwal)
        ch3_tbl_idx += 1

    # =========================================================================
    # DAFTAR REFERENSI
    # =========================================================================
    # Format Harvard style, unnumbered level 1 heading, sorted alphabetically
    ref_heading_style = 'Reference Heading' if 'Reference Heading' in doc.styles else 'Heading 1'
    doc.add_paragraph("DAFTAR REFERENSI", style=ref_heading_style)

    sorted_refs = sorted(daftar_referensi, key=lambda r: r.lower().strip())
    for r in sorted_refs:
        add_ref(r)

    # =========================================================================
    # LAMPIRAN (OPSIONAL)
    # =========================================================================
    if lampiran_data:
        for lamp in lampiran_data:
            lamp_title = clean_heading_title(lamp.get("title", ""), level=1)
            lamp_style = 'Appendix Heading 1' if 'Appendix Heading 1' in doc.styles else 'Heading 1'
            doc.add_paragraph(lamp_title, style=lamp_style)
            for sub in lamp.get("sections", []):
                sub_title = clean_heading_title(sub.get("title", ""), level=2)
                sub_style = 'Appendix Heading 2' if 'Appendix Heading 2' in doc.styles else 'Heading 2'
                doc.add_paragraph(sub_title, style=sub_style)
                for p_text in sub.get("paragraphs", []):
                    add_p(p_text)

    # =========================================================================
    # WORD SETTINGS: UPDATE FIELDS OTOMATIS
    # =========================================================================
    # Mengaktifkan w:updateFields agar Word secara otomatis memperbarui
    # DAFTAR ISI, DAFTAR TABEL, DAFTAR GAMBAR, dan nomor halaman saat dibuka.
    if doc.settings.element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields') is None:
        update_fields = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
        doc.settings.element.append(update_fields)

    doc.save(output_path)
    return {
        "status": "SUCCESS",
        "output_path": output_path,
        "judul": metadata.get("judul"),
        "penulis": metadata.get("nama_mahasiswa"),
        "file_size_bytes": os.path.getsize(output_path)
    }

def inspect_doc(docx_path: str) -> Dict[str, Any]:
    """
    Memeriksa struktur dokumen proposal skripsi (.docx) di dalam workspace.
    Mengembalikan informasi lengkap: paragraf, tabel, seksi, perkiraan kata,
    daftar heading bab/subbab, caption tabel/gambar, dan rincian tabel.
    """
    if not os.path.exists(docx_path):
        return {"error": f"File {docx_path} tidak ditemukan."}

    doc = docx.Document(docx_path)
    headings = []
    captions = []
    tables_info = []
    total_words = 0

    for p in doc.paragraphs:
        total_words += len(p.text.split())
        if p.style and (p.style.name.startswith("Heading") or p.style.name in ["Reference Heading", "Appendix Heading 1", "Appendix Heading 2"]):
            headings.append({"style": p.style.name, "text": p.text.strip()})
        if p.style and p.style.name == "Caption" and p.text.strip():
            captions.append(p.text.strip())

    for idx, t in enumerate(doc.tables):
        tables_info.append({
            "table_index": idx + 1,
            "rows": len(t.rows),
            "columns": len(t.columns)
        })

    return {
        "file": os.path.basename(docx_path),
        "total_paragraphs": len(doc.paragraphs),
        "total_tables": len(doc.tables),
        "total_sections": len(doc.sections),
        "approx_words": total_words,
        "headings": headings,
        "captions": captions,
        "tables_info": tables_info
    }

def record_version_change(workspace_dir: str, from_version: str, to_version: str, notes: str) -> str:
    log_file = os.path.join(workspace_dir, "version_history.json")
    history = []
    if os.path.exists(log_file):
        try:
            with open(log_file, "r", encoding="utf-8") as f:
                history = json.load(f)
        except Exception:
            history = []

    history.append({
        "timestamp": datetime.now().isoformat(),
        "from_version": from_version,
        "to_version": to_version,
        "notes": notes
    })

    with open(log_file, "w", encoding="utf-8") as f:
        json.dump(history, f, indent=2, ensure_ascii=False)

    return f"Versi {to_version} berhasil dicatat di version_history.json."
